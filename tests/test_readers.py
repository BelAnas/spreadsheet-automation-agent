import http.server
import threading
from functools import partial

import pytest
from openpyxl import Workbook

from sheetagent.readers import ReadError, collect, google_sheet_csv_url, read_source, sniff_delimiter


def test_delimiter_is_sniffed():
    assert sniff_delimiter("a;b;c\n1;2;3\n4;5;6") == ";"
    assert sniff_delimiter("a\tb\n1\t2") == "\t"
    assert sniff_delimiter("a|b|c\n1|2|3") == "|"
    # decimal commas inside a semicolon file must not win
    assert sniff_delimiter("Date;Price;Amount\n01.09.2026;12,50;25,00\n02.09.2026;3,10;6,20") == ";"


def test_csv_in_windows_encoding_and_with_bom(tmp_path):
    (tmp_path / "a.csv").write_bytes("Client;Prix\nCafé Müller;12,50 €\n".encode("cp1252"))
    (tmp_path / "b.csv").write_bytes("Region,Target\nNorth,10\n".encode("utf-8-sig"))
    a = read_source(tmp_path / "a.csv")[0].grid
    assert a.iloc[1, 0] == "Café Müller" and a.iloc[1, 1] == "12,50 €"
    assert read_source(tmp_path / "b.csv")[0].grid.iloc[0, 0] == "Region"  # no byte-order mark left


def test_workbook_gives_one_grid_per_visible_sheet(tmp_path):
    wb = Workbook()
    wb.active.title = "Data"
    wb.active.append(["Date", "Amount"])
    wb.active.append(["2026-09-01", 5])
    hidden = wb.create_sheet("Hidden")
    hidden.append(["x"])
    hidden.sheet_state = "hidden"
    wb.create_sheet("Empty")
    wb.save(tmp_path / "book.xlsx")
    tables = read_source(tmp_path / "book.xlsx")
    assert [t.sheet for t in tables] == ["Data"]
    assert tables[0].label == "book.xlsx [Data]"


def test_broken_and_unsupported_files_raise_a_readable_error(tmp_path):
    (tmp_path / "broken.xlsx").write_text("this is not a workbook")
    (tmp_path / "notes.pdf").write_text("x")
    with pytest.raises(ReadError, match="cannot open"):
        read_source(tmp_path / "broken.xlsx")
    with pytest.raises(ReadError, match="unsupported"):
        read_source(tmp_path / "notes.pdf")


def test_collect_expands_folders_and_skips_lock_files(tmp_path):
    for name in ("b.csv", "a.xlsx", "~$a.xlsx", "readme.md", ".hidden.csv"):
        (tmp_path / name).write_text("x")
    found = collect([tmp_path, "https://example.com/x.csv"])
    assert [getattr(p, "name", p) for p in found] == ["a.xlsx", "b.csv", "https://example.com/x.csv"]


def test_google_sheets_link_becomes_a_csv_export_link():
    url = "https://docs.google.com/spreadsheets/d/1AbC_dEf-123/edit#gid=42"
    assert google_sheet_csv_url(url) == "https://docs.google.com/spreadsheets/d/1AbC_dEf-123/export?format=csv&gid=42"
    assert google_sheet_csv_url("https://docs.google.com/spreadsheets/d/1AbC/edit?usp=sharing").endswith("/export?format=csv")
    assert google_sheet_csv_url("https://example.com/t.csv") == "https://example.com/t.csv"


@pytest.fixture
def web_folder(tmp_path):
    """A local web server, standing in for the CSV export of a shared Google Sheet."""
    handler = partial(http.server.SimpleHTTPRequestHandler, directory=str(tmp_path))
    handler.log_message = lambda *a, **k: None
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield tmp_path, f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


def test_table_is_read_from_a_link(web_folder):
    folder, base = web_folder
    (folder / "targets.csv").write_text("Region,Month,Target\nNorth,2026-09,1000\n", encoding="utf-8")
    table = read_source(f"{base}/targets.csv")[0]
    assert table.source == "targets.csv" and table.grid.shape == (2, 3)


def test_dead_link_raises_a_readable_error(web_folder):
    with pytest.raises(ReadError, match="download failed"):
        read_source(f"{web_folder[1]}/missing.csv")
