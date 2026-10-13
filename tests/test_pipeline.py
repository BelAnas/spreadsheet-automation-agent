import json

import pytest
from openpyxl import load_workbook

from evals.scoring import score
from sheetagent import metrics
from sheetagent.agent import Agent
from sheetagent.cli import main, run_report
from sheetagent.report import publish

from .conftest import ScriptedLLM


@pytest.mark.parametrize("period", ["2026-08", "2026-09"])
def test_demo_month_is_recovered_exactly(demo, period):
    root, truths = demo
    run = Agent().clean([root / "inbox" / period, root / "reference"], period)
    result = score(run.sales, truths[period])
    assert result["row_recovery"] == 1.0 and result["row_precision"] == 1.0
    assert run.sales["amount"].sum() == pytest.approx(truths[period]["amount"].sum())
    merged = run.sales.merge(truths[period][["order_id", "sku", "region"]], on=["order_id", "sku"], suffixes=("", "_true"))
    assert (merged["region"] == merged["region_true"]).all()
    reasons = {r["reason"].split(":")[0] for r in run.rejected}
    assert reasons == {"no product code", "date not understood", "no usable amount", "duplicate line",
                       f"outside the report month {period}"}
    assert any("not in the product list" in w for w in run.warnings)


def test_same_input_gives_the_same_report(demo):
    root, _ = demo
    inputs = [root / "inbox" / "2026-09", root / "reference"]
    a, b = (Agent().report(Agent().clean(inputs, "2026-09"), {}) for _ in range(2))
    assert a.facts == b.facts and a.summary.text == b.summary.text


def test_report_month_is_found_when_not_given(demo):
    root, _ = demo
    assert Agent().clean([root / "inbox" / "2026-08", root / "reference"]).period == "2026-08"


def test_outputs_and_month_over_month(demo, tmp_path):
    root, _ = demo
    for period in ("2026-08", "2026-09"):
        run = run_report([str(root / "inbox" / period), str(root / "reference")], period, tmp_path, quiet=True)
    assert run.facts["previous"]["period"] == "2026-08" and run.facts["change"]["revenue_pct"] == 6.2
    assert set(json.loads((tmp_path / "history.json").read_text())) == {"2026-08", "2026-09"}

    wb = load_workbook(run.outputs["excel"])
    assert wb.sheetnames == ["Summary", "By region", "By category", "Top products", "Daily", "Clean data",
                             "Rejected rows", "Column mapping", "Audit log"]
    assert len(wb["Summary"]._charts) == 3
    assert wb["Clean data"].max_row == len(run.sales) + 1
    assert wb["Rejected rows"].max_row == len(run.rejected) + 1
    assert wb["Summary"]["A5"].value == run.summary.text

    html = open(run.outputs["html"], encoding="utf-8").read()
    assert run.summary.text.replace("&", "&amp;") in html and "Plotly.react" in html and "Below target" in html
    log = json.loads(open(run.outputs["log"]).read())
    assert log["facts"]["totals"]["revenue"] == run.facts["totals"]["revenue"] and len(log["trace"]) == 7


def test_cdn_page_is_small(demo, tmp_path):
    root, _ = demo
    run = run_report([str(root / "inbox" / "2026-09"), str(root / "reference")], "2026-09", tmp_path, plotly="cdn", quiet=True)
    assert (tmp_path / "2026-09" / "sales_report_2026-09.html").stat().st_size < 200_000
    assert "</script><script>alert" not in open(run.outputs["html"]).read()


def test_text_from_files_is_escaped_in_the_page(demo, tmp_path):
    root, _ = demo
    run = Agent().report(Agent().clean([root / "inbox" / "2026-09", root / "reference"], "2026-09"), {})
    run.facts["regions"][0]["region"] = "<script>alert(1)</script>"
    run.rejected[0]["original"] = "</script><img src=x onerror=alert(1)>"
    publish(run, tmp_path, plotly="cdn")
    html = open(run.outputs["html"], encoding="utf-8").read()
    assert "<script>alert(1)</script>" not in html and "<img src=x" not in html


def test_broken_file_is_reported_and_the_rest_still_runs(demo, tmp_path):
    root, _ = demo
    (tmp_path / "broken.xlsx").write_text("not a workbook")
    (tmp_path / "notes.csv").write_text("Topic,Count\nmeetings,3\ncalls,5\n")
    run = Agent().clean([root / "inbox" / "2026-09", root / "reference", tmp_path], "2026-09")
    assert len(run.sales) == 862
    assert any("broken.xlsx" in w and "cannot open" in w for w in run.warnings)
    assert any("notes.csv" in w and "not recognised" in w for w in run.warnings)


def test_no_product_list_and_no_targets_still_gives_a_report(demo):
    root, truths = demo
    run = Agent().report(Agent().clean([root / "inbox" / "2026-09"], "2026-09"), {})
    assert run.ok and run.facts["target"] is None and run.facts["totals"]["margin_pct"] is None
    assert any("No product list" in w for w in run.warnings)
    # Without the product list, codes cannot be repaired, but the revenue is still complete.
    assert run.facts["totals"]["revenue"] == pytest.approx(truths["2026-09"]["amount"].sum(), abs=0.01)


def test_empty_folder_gives_no_report(tmp_path):
    run = Agent().report(Agent().clean([tmp_path]), {})
    assert not run.ok and run.sales.empty
    assert main(["run", str(tmp_path), "--out", str(tmp_path / "out")]) == 1


def test_model_decisions_are_remembered_for_the_next_run(demo, tmp_path):
    root, _ = demo
    inputs = [root / "inbox" / "2026-09", root / "reference"]
    memory = tmp_path / "mapping_memory.json"
    first = ScriptedLLM(rule=lambda prompt: "customer" if "'Store'" in prompt else "none")
    run = Agent(first, memory_path=memory).clean(inputs, "2026-09")
    assert first.usage.calls >= 1 and json.loads(memory.read_text()) == {"sales": {"store": "customer"}}
    assert run.sales.loc[run.sales["region"] == "West", "customer"].notna().all()
    second = ScriptedLLM(rule=lambda prompt: "none")
    again = Agent(second, memory_path=memory).clean(inputs, "2026-09")
    assert not any("'Store'" in p for p in second.prompts)
    assert any(m["column"] == "Store" and m["method"] == "memory" for m in again.mappings)


def test_metrics_on_a_tiny_table():
    import pandas as pd
    from datetime import date
    sales = pd.DataFrame({"order_id": ["a", "a", "b", None], "order_date": [date(2026, 9, 1)] * 3 + [date(2026, 9, 2)],
                          "region": ["North", "North", "South", "South"], "sku": ["x", "y", "x", "x"],
                          "quantity": [1.0, 2.0, -1.0, 3.0], "amount": [10.0, 30.0, -10.0, 30.0]})
    products = pd.DataFrame({"sku": ["x"], "product_name": ["X"], "category": ["Tools"], "unit_cost": [4.0]})
    targets = pd.DataFrame({"region": ["North"], "month": ["2026-09"], "target_revenue": [50.0]})
    facts = metrics.compute(metrics.enrich(sales, products), targets, "2026-09", {"2026-08": {
        "revenue": 50.0, "units": 5.0, "orders": 2, "regions": {"North": 20.0}}})
    t = facts["totals"]
    assert (t["revenue"], t["orders"], t["lines"], t["returns_value"], t["return_lines"]) == (60.0, 2, 4, 10.0, 1)
    assert t["gross_margin"] == 18.0 and t["margin_pct"] == 60.0  # only lines with a known cost count
    assert facts["change"]["revenue_pct"] == 20.0
    north = next(r for r in facts["regions"] if r["region"] == "North")
    assert (north["attainment_pct"], north["change_pct"], north["share_pct"]) == (80.0, 100.0, 66.7)
    assert {c["category"] for c in facts["categories"]} == {"Tools", metrics.UNKNOWN_CATEGORY}
    assert facts["best_day"] == {"date": "2026-09-01", "revenue": 30.0}
    assert metrics.previous_period("2026-01") == "2025-12"


def test_command_line_demo(tmp_path, capsys):
    assert main(["demo", "--data", str(tmp_path / "data"), "--out", str(tmp_path / "reports"), "--plotly", "cdn"]) == 0
    out = capsys.readouterr().out
    assert "September 2026" in out and "+6.2% on August 2026" in out and "Excel report" in out
    assert (tmp_path / "reports" / "2026-09" / "sales_report_2026-09.xlsx").exists()


def test_openai_provider_without_key_is_explained(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("OPENAI_BASE_URL", raising=False)
    assert main(["run", str(tmp_path), "--provider", "openai", "--out", str(tmp_path / "o")]) == 2
    assert "OPENAI_API_KEY" in capsys.readouterr().err


def test_mixed_currencies_are_flagged_not_converted(tmp_path):
    (tmp_path / "a.csv").write_text("Date,SKU,Qty,Amount\n2026-09-01,HT-0101,1,$10.00\n2026-09-02,HT-0102,2,$20.00\n")
    (tmp_path / "b.csv").write_text("Date;SKU;Qty;Amount\n2026-09-01;HT-0101;1;10,00 €\n2026-09-02;HT-0102;2;20,00 €\n")
    run = Agent().clean([tmp_path], "2026-09")
    assert run.sales["amount"].sum() == 60.0
    assert any("different currency marks" in w for w in run.warnings)


def test_template_writes_the_summary_unless_the_model_is_asked_for(demo):
    root, _ = demo
    inputs = [root / "inbox" / "2026-09", root / "reference"]
    text = ("September 2026 brought revenue of 253,055 USD across the 3 regions, a good month for the company "
            "with North in the lead.")
    quiet = ScriptedLLM(rule=lambda prompt: text if "Facts:" in prompt else "none")
    default = Agent(quiet).report(Agent(quiet).clean(inputs, "2026-09"), {})
    assert default.summary.source == "template" and not any("Facts:" in p for p in quiet.prompts)
    asked = Agent(quiet, summary_writer="model")
    assert asked.report(asked.clean(inputs, "2026-09"), {}).summary.source == "model"
