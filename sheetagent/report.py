"""Publish a run: the Excel workbook, the web page, the summary text and the run log."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from jinja2 import Environment, PackageLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from plotly.offline import get_plotlyjs
from plotly.utils import PlotlyJSONEncoder

from .agent import Run
from .charts import figures
from .summary import money, pct, signed

PLOTLY_CDN = '<script src="https://cdn.plot.ly/plotly-2.35.2.min.js" charset="utf-8"></script>'
_HEAD_FILL = PatternFill("solid", fgColor="E7E6E2")
_BLUE, _INK = "2A78D6", "52514E"

CLEAN_COLUMNS = ["order_date", "region", "order_id", "customer", "sku", "product_name", "category", "quantity",
                 "unit_price", "amount", "source", "row"]


def tiles(facts: dict) -> list[dict]:
    """The headline figures, each with the comparison that gives it meaning."""
    t, cur, change = facts["totals"], facts["currency"], facts["change"] or {}
    last = f"against {facts['previous']['label']}" if facts["previous"] else ""

    def delta(key: str) -> str:
        value = change.get(key)
        if value is None:
            return "no previous month on file"
        return f"{'▲' if value >= 0 else '▼'} {abs(value):.1f}% {last}"

    out = [{"label": f"Revenue ({cur})", "value": money(t["revenue"]), "note": delta("revenue_pct")},
           {"label": "Orders", "value": f"{t['orders']:,}", "note": delta("orders_pct")},
           {"label": "Units", "value": f"{t['units']:,.0f}", "note": delta("units_pct")},
           {"label": f"Average order ({cur})", "value": money(t["average_order_value"]), "note": f"{t['lines']:,} order lines"}]
    if t["margin_pct"] is not None:
        out.append({"label": "Gross margin", "value": pct(t["margin_pct"]), "note": f"{money(t['gross_margin'])} {cur}"})
    if facts["target"]:
        out.append({"label": "Of target", "value": pct(facts["target"]["attainment_pct"]),
                    "note": f"target {money(facts['target']['total'])} {cur}"})
    return out


# ----------------------------------------------------------------------------- Excel

def _sheet(wb: Workbook, title: str, frame: pd.DataFrame, money_columns: tuple[str, ...] = ()) -> object:
    ws = wb.create_sheet(title)
    ws.append([str(c).replace("_", " ").capitalize() for c in frame.columns])
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), _HEAD_FILL
    for row in frame.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    for i, column in enumerate(frame.columns, start=1):
        longest = max([len(str(column))] + [len(str(v)) for v in frame[column].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
        if column in money_columns:
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    return ws


def _style(chart, title: str, width: float = 16, height: float = 7.5) -> None:
    chart.title, chart.width, chart.height = title, width, height
    chart.legend.position = "b"
    # openpyxl hides the axes of new charts unless this is set.
    chart.x_axis.delete = chart.y_axis.delete = False


def write_excel(run: Run, path: Path, title: str) -> None:
    facts = run.facts
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A2"] = title, f"{facts['period_label']} - all amounts in {facts['currency']}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A4"] = "Summary"
    ws["A4"].font = Font(bold=True)
    ws.merge_cells("A5:H11")
    ws["A5"] = run.summary.text
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A12"] = f"Written by: {run.summary.source}. Figures computed by code from the Clean data sheet."
    ws["A12"].font = Font(italic=True, color=_INK)
    ws["A14"] = "Key figures"
    ws["A14"].font = Font(bold=True)
    for offset, tile in enumerate(tiles(facts)):
        ws.cell(15 + offset, 1, tile["label"])
        ws.cell(15 + offset, 3, tile["value"]).alignment = Alignment(horizontal="right")
        ws.cell(15 + offset, 4, tile["note"])
    ws.column_dimensions["A"].width = 24

    regions = pd.DataFrame(facts["regions"]).rename(columns={
        "share_pct": "share_%", "attainment_pct": "of_target_%", "change_pct": "change_vs_last_month_%"})
    r_ws = _sheet(wb, "By region", regions, ("revenue", "target", "previous_revenue"))
    c_ws = _sheet(wb, "By category", pd.DataFrame(facts["categories"]).rename(columns={
        "share_pct": "share_%", "margin_pct": "margin_%"}), ("revenue",))
    _sheet(wb, "Top products", pd.DataFrame(facts["top_products"]), ("revenue",))
    d_ws = _sheet(wb, "Daily", pd.DataFrame(facts["daily"]), ("revenue",))

    clean = run.sales[CLEAN_COLUMNS].sort_values(["order_date", "source", "row"])
    clean_ws = _sheet(wb, "Clean data", clean, ("unit_price", "amount"))
    clean_ws.auto_filter.ref = clean_ws.dimensions
    _sheet(wb, "Rejected rows", pd.DataFrame(run.rejected, columns=["source", "row", "reason", "original"]))
    _sheet(wb, "Column mapping", pd.DataFrame(run.mappings, columns=["source", "table", "column", "field", "method",
                                                                     "confidence", "note"]))
    _sheet(wb, "Audit log", pd.DataFrame(run.audit, columns=["step", "source", "action", "count", "example"]))

    # Native Excel charts: they stay editable, and they update if someone corrects a figure.
    n = len(facts["regions"])
    bar = BarChart()
    columns = list(regions.columns)
    for name, colour in (("revenue", _BLUE), ("target", _INK)):
        if name == "target" and regions["target"].isna().all():
            continue
        col = columns.index(name) + 1
        bar.add_data(Reference(r_ws, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        bar.series[-1].graphicalProperties.solidFill = colour
        bar.series[-1].graphicalProperties.line.solidFill = colour
    bar.set_categories(Reference(r_ws, min_col=1, min_row=2, max_row=n + 1))
    bar.gapWidth = 80
    _style(bar, "Revenue by region against target")
    ws.add_chart(bar, "J2")

    line = LineChart()
    line.add_data(Reference(d_ws, min_col=2, min_row=1, max_row=len(facts["daily"]) + 1), titles_from_data=True)
    line.set_categories(Reference(d_ws, min_col=1, min_row=2, max_row=len(facts["daily"]) + 1))
    line.series[0].graphicalProperties.line.solidFill = _BLUE
    line.series[0].graphicalProperties.line.width = 20000  # EMU, about 1.6 pt
    line.series[0].smooth = False
    _style(line, "Revenue per day", width=24)
    line.legend = None
    ws.add_chart(line, "A24")

    cat = BarChart()
    cat.type = "bar"  # horizontal, so the category names are readable
    k = len(facts["categories"])
    cat.add_data(Reference(c_ws, min_col=2, min_row=1, max_row=k + 1), titles_from_data=True)
    cat.set_categories(Reference(c_ws, min_col=1, min_row=2, max_row=k + 1))
    cat.series[0].graphicalProperties.solidFill = _BLUE
    cat.y_axis.scaling.orientation = "minMax"
    cat.x_axis.scaling.orientation = "maxMin"  # largest category at the top
    _style(cat, "Revenue by category")
    cat.legend = None
    ws.add_chart(cat, "J18")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------------- web page

def render_html(run: Run, title: str, plotly: str = "inline") -> str:
    env = Environment(loader=PackageLoader("sheetagent", "templates"), autoescape=select_autoescape(["html", "j2"]))
    # "</" inside a script block would end it, so it is escaped in the embedded data.
    figures_json = json.dumps(figures(run.facts), cls=PlotlyJSONEncoder).replace("</", "<\\/")
    script = PLOTLY_CDN if plotly == "cdn" else f"<script>{get_plotlyjs()}</script>"
    return env.get_template("report.html.j2").render(
        title=title, facts=run.facts, q=run.facts["quality"], summary=run.summary, tiles=tiles(run.facts),
        trace=run.trace, mappings=run.mappings, rejected=run.rejected[:50], llm=run.llm, seconds=run.seconds,
        generated=datetime.now().strftime("%d %b %Y %H:%M"), figures_json=figures_json, plotly_script=script,
        money=money, pct=pct, signed=signed)


def publish(run: Run, out_dir: Path, title: str = "Monthly sales report", plotly: str = "inline") -> dict[str, str]:
    """Write every output of a run into `out_dir/<period>/` and return the paths."""
    folder = out_dir / run.period
    folder.mkdir(parents=True, exist_ok=True)
    paths = {"excel": folder / f"sales_report_{run.period}.xlsx", "html": folder / f"sales_report_{run.period}.html",
             "summary": folder / "summary.md", "log": folder / "run.json"}
    write_excel(run, paths["excel"], title)
    paths["html"].write_text(render_html(run, title, plotly), encoding="utf-8")
    paths["summary"].write_text(f"# {title} - {run.facts['period_label']}\n\n{run.summary.text}\n\n"
                                f"_Written by: {run.summary.source}._\n", encoding="utf-8")
    log = {"period": run.period, "facts": run.facts, "summary": {"text": run.summary.text, "source": run.summary.source,
                                                                 "attempts": run.summary.attempts,
                                                                 "rejected_numbers": run.summary.rejected_numbers},
           "trace": run.trace, "mappings": run.mappings, "audit": run.audit, "warnings": run.warnings,
           "rejected": len(run.rejected), "llm": run.llm, "seconds": round(run.seconds, 3)}
    paths["log"].write_text(json.dumps(log, indent=2, default=str), encoding="utf-8")
    run.outputs = {k: str(v) for k, v in paths.items()}
    return run.outputs
