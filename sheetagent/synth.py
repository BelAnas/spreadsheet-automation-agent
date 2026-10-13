"""Synthetic data: a clean truth, and messy spreadsheet files made from it.

Everything here is fictional (a garden-supplies wholesaler with three regions). The clean
truth is kept so that the cleaning pipeline can be scored against a known correct answer,
which is impossible with real client files.
"""
from __future__ import annotations

import calendar
import csv
import random
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

REGIONS = ["North", "South", "West"]

CATEGORIES: dict[str, tuple[str, list[str]]] = {
    "Hand tools": ("HT", ["Steel trowel", "Hand fork", "Bypass pruner", "Hedge shears", "Weeding knife",
                          "Folding saw", "Bulb planter", "Soil scoop"]),
    "Watering": ("WT", ["Watering can 10 L", "Garden hose 25 m", "Hose reel", "Spray nozzle",
                        "Drip kit 20 pots", "Tap connector", "Rain barrel 200 L", "Soaker hose 15 m"]),
    "Soil and compost": ("SC", ["Potting mix 40 L", "Peat-free compost 50 L", "Bark mulch 60 L",
                                "Perlite 10 L", "Seed compost 20 L", "Topsoil 25 kg",
                                "Tomato feed 1 L", "Lawn feed 5 kg"]),
    "Planters": ("PL", ["Terracotta pot 20 cm", "Terracotta pot 30 cm", "Window box 60 cm",
                        "Hanging basket 35 cm", "Raised bed kit", "Trough planter 80 cm",
                        "Plant saucer 30 cm", "Seed tray set"]),
    "Seeds": ("SD", ["Tomato seeds", "Basil seeds", "Carrot seeds", "Lettuce mix seeds",
                     "Sunflower seeds", "Wildflower mix 100 g", "Lawn seed 1 kg", "Pea seeds"]),
    "Outdoor lighting": ("OL", ["Solar stake light", "Solar string lights 10 m", "Path light set of 4",
                                "Wall lantern", "Deck light set of 6", "Motion flood light",
                                "Spike spotlight", "Festoon lights 15 m"]),
}

# Header spellings used when files are made. "seen" spellings are in the alias lists of
# schema.py; "unseen" ones are deliberately absent from them, so the benchmark can report
# how the pipeline does on names nobody prepared it for.
HEADER_POOL: dict[str, dict[str, list[str]]] = {
    "order_id": {"seen": ["Order ID", "Order No.", "Invoice #"],
                 "unseen": ["Doc Ref", "Ref.", "Transaction", "Bon de commande", "Ticket no"]},
    "order_date": {"seen": ["Order Date", "Date", "Invoice Date"],
                   "unseen": ["Dt", "Posted on", "Booked", "Date de commande", "Day"]},
    "sku": {"seen": ["SKU", "Product Code", "Item No."],
            "unseen": ["Part #", "Ref article", "Material", "Code"]},
    "quantity": {"seen": ["Quantity", "Qty", "Units"],
                 "unseen": ["Qté", "Pieces", "Nb", "Vol.", "Count"]},
    "unit_price": {"seen": ["Unit Price", "Price"],
                   "unseen": ["PU", "Rate", "Each", "Prix unitaire", "Unit $"]},
    "amount": {"seen": ["Amount", "Total", "Line Total", "Revenue"],
               "unseen": ["Ext. Price", "Value", "Montant", "Net", "Turnover"]},
    "customer": {"seen": ["Customer", "Client", "Account"],
                 "unseen": ["Buyer", "Sold to", "Store", "Outlet"]},
    "region": {"seen": ["Region", "Branch"], "unseen": ["Area", "Zone", "Depot"]},
}

JUNK_COLUMNS = {"Sales rep": ["A. Reyes", "M. Okafor", "L. Brandt", "S. Imai"],
                "Notes": ["", "", "", "call back", "pallet", "urgent"],
                "Warehouse": ["W1", "W2", "W3"]}

DATE_STYLES = ["iso", "dmy_slash", "mdy_slash", "dmy_dot", "text_mon", "mon_text", "native", "dmy_short"]
NUMBER_STYLES = ["plain", "us_currency", "eu", "eu_space", "code_suffix"]
_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ----------------------------------------------------------------------------- clean truth

def make_products(seed: int = 7) -> pd.DataFrame:
    rng = random.Random(seed)
    rows = []
    for c_index, (category, (prefix, names)) in enumerate(CATEGORIES.items()):
        for p_index, name in enumerate(names):
            cost = round(rng.uniform(1.5, 38.0), 2)
            rows.append({"sku": f"{prefix}-{(c_index + 1) * 100 + p_index + 1:04d}", "product_name": name,
                         "category": category, "unit_cost": cost,
                         "list_price": round(cost * rng.uniform(1.35, 1.9), 2)})
    return pd.DataFrame(rows)


def make_customers(seed: int = 7) -> list[str]:
    rng = random.Random(seed + 1)
    first = ["Greenleaf", "Oakfield", "Riverside", "Hilltop", "Meadow", "Willow", "Cedar", "Stonebridge",
             "Bluebell", "Ashgrove", "Fernhill", "Maple"]
    last = ["Garden Centre", "Nursery", "Home and Garden", "Landscapes", "Growers", "Garden Supplies"]
    names = [f"{a} {b}" for a in first for b in last]
    rng.shuffle(names)
    return names[:36]


def make_sales(period: str, products: pd.DataFrame, seed: int = 7,
               orders_per_region: dict[str, int] | None = None) -> pd.DataFrame:
    """One month of order lines for every region. `period` is 'YYYY-MM'."""
    year, month = (int(x) for x in period.split("-"))
    rng = random.Random(f"{seed}-{period}")
    customers = make_customers(seed)
    orders_per_region = orders_per_region or {"North": 150, "South": 115, "West": 135}
    days = calendar.monthrange(year, month)[1]
    rows = []
    for region in REGIONS:
        for n in range(orders_per_region[region]):
            # Fewer orders on weekends, like a trade wholesaler.
            while True:
                day = date(year, month, rng.randint(1, days))
                if day.weekday() < 5 or rng.random() < 0.25:
                    break
            order_id = f"{region[0]}{year % 100:02d}{month:02d}-{n + 1:04d}"
            customer = rng.choice(customers)
            # One line per product inside an order, as an ERP writes it.
            for index in rng.sample(range(len(products)), rng.choice([1, 1, 2, 2, 3, 4])):
                p = products.iloc[index]
                qty = rng.choice([1, 2, 3, 4, 5, 6, 8, 10, 12, 20, 24, 40])
                if rng.random() < 0.02:
                    qty = -rng.choice([1, 2, 3])  # a return
                price = round(p["list_price"] * (1 - rng.choice([0, 0, 0, 0.05, 0.10])), 2)
                rows.append({"order_id": order_id, "order_date": day, "region": region, "sku": p["sku"],
                             "quantity": qty, "unit_price": price, "amount": round(qty * price, 2),
                             "customer": customer})
    return pd.DataFrame(rows)


def make_targets(periods: list[str], seed: int = 7) -> pd.DataFrame:
    rng = random.Random(seed + 2)
    base = {"North": 92000, "South": 84000, "West": 76000}
    return pd.DataFrame([{"region": r, "month": p, "target_revenue": int(base[r] * rng.uniform(0.95, 1.08) // 500 * 500)}
                         for p in periods for r in REGIONS])


# ----------------------------------------------------------------------------- the mess

@dataclass
class MessSpec:
    """How one sales file is spoiled. Every option is something seen in real exports."""
    fmt: str = "xlsx"  # xlsx | csv
    headers: dict[str, str] = field(default_factory=dict)  # field -> header text
    order: list[str] = field(default_factory=list)  # field names (and junk names) left to right
    banner_rows: int = 0
    date_style: str = "iso"
    mixed_native_dates: bool = False  # some date cells are real Excel dates, the rest text
    number_style: str = "plain"
    qty_suffix: bool = False  # "12 pcs"
    sku_mess: float = 0.0  # share of product codes written badly
    blank_rows: int = 0
    totals_row: bool = False
    repeat_header: bool = False
    duplicates: int = 0
    bad_rows: int = 0
    split_sheets: int = 1
    extra_sheet: bool = False
    delimiter: str = ","
    encoding: str = "utf-8"


def format_date(d: date, style: str) -> object:
    if style == "iso":
        return d.isoformat()
    if style == "dmy_slash":
        return d.strftime("%d/%m/%Y")
    if style == "mdy_slash":
        return d.strftime("%m/%d/%Y")
    if style == "dmy_dot":
        return d.strftime("%d.%m.%Y")
    if style == "dmy_short":
        return d.strftime("%d/%m/%y")
    if style == "text_mon":
        return f"{d.day} {_MONTHS[d.month - 1]} {d.year}"
    if style == "mon_text":
        return f"{_MONTHS[d.month - 1]} {d.day}, {d.year}"
    return datetime(d.year, d.month, d.day)  # native


def format_number(x: float, style: str) -> object:
    if style == "plain":
        return x
    whole = f"{abs(x):,.2f}"  # 1,234.50
    if style == "us_currency":
        return f"(${whole})" if x < 0 else f"${whole}"
    if style == "code_suffix":
        return f"{x:.2f} USD"
    eu = whole.replace(",", "§").replace(".", ",")
    sign = "-" if x < 0 else ""
    if style == "eu":
        return sign + eu.replace("§", ".")
    return sign + eu.replace("§", " ") + " €"  # eu_space


def mess_sku(sku: str, rng: random.Random) -> str:
    prefix, digits = sku.split("-")
    return rng.choice([sku.lower(), f" {sku} ", f"{prefix}_{digits}", f"{prefix}{digits}",
                       f"{prefix} {digits}", f"{prefix}-{digits.replace('0', 'O', 1)}",
                       f"{prefix.lower()}-{int(digits)}"])


def random_spec(rng: random.Random, unseen_share: float = 0.4) -> MessSpec:
    """A random but plausible file layout for the benchmark."""
    fmt = rng.choice(["xlsx", "xlsx", "csv"])
    fields = ["order_id", "order_date", "sku", "quantity", "unit_price", "amount", "customer", "region"]
    # Real exports often carry only two of quantity, price and total; the third is derived.
    dropped = rng.choice([None, None, "unit_price", "amount"])
    if dropped:
        fields.remove(dropped)
    if rng.random() < 0.5:
        fields.remove("region")  # the region is then only in the file name
    if rng.random() < 0.2:
        fields.remove("customer")
    headers = {f: rng.choice(HEADER_POOL[f]["unseen" if rng.random() < unseen_share else "seen"]) for f in fields}
    order = fields + rng.sample(list(JUNK_COLUMNS), k=rng.choice([0, 1, 1, 2]))
    rng.shuffle(order)
    date_style = rng.choice(DATE_STYLES)
    if fmt == "csv" and date_style == "native":
        date_style = "iso"
    number_style = rng.choice(NUMBER_STYLES)
    delimiter = rng.choice([",", ";", "\t", "|"]) if fmt == "csv" else ","
    return MessSpec(
        fmt=fmt, headers=headers, order=order, banner_rows=rng.choice([0, 0, 1, 2, 3, 4]),
        date_style=date_style, mixed_native_dates=fmt == "xlsx" and date_style != "native" and rng.random() < 0.3,
        number_style=number_style, qty_suffix=rng.random() < 0.25, sku_mess=rng.choice([0.0, 0.1, 0.3]),
        blank_rows=rng.choice([0, 0, 2, 5]), totals_row=rng.random() < 0.5, repeat_header=rng.random() < 0.3,
        duplicates=rng.choice([0, 0, 3, 8]), bad_rows=rng.choice([0, 2, 5]),
        split_sheets=rng.choice([1, 1, 1, 3]) if fmt == "xlsx" else 1, extra_sheet=fmt == "xlsx" and rng.random() < 0.4,
        delimiter=delimiter, encoding=rng.choice(["utf-8", "utf-8-sig", "cp1252"]) if fmt == "csv" else "utf-8",
    )


def _cell(row: dict, name: str, spec: MessSpec, rng: random.Random) -> object:
    if name in JUNK_COLUMNS:
        return rng.choice(JUNK_COLUMNS[name])
    value = row[name]
    if name == "order_date":
        if spec.mixed_native_dates and rng.random() < 0.4:
            return format_date(value, "native")
        return format_date(value, spec.date_style)
    if name == "quantity":
        return f"{value} pcs" if spec.qty_suffix and rng.random() < 0.6 else value
    if name in ("unit_price", "amount"):
        return format_number(value, spec.number_style)
    if name == "sku" and rng.random() < spec.sku_mess:
        return mess_sku(value, rng)
    return value


def _bad_row(spec: MessSpec, sample: dict, n: int, rng: random.Random) -> list[object]:
    """A line no cleaning can save: the code, the date or both numbers are unusable."""
    row = dict(sample)
    row["order_id"] = f"ERR-{n:03d}"
    cells = {name: _cell(row, name, spec, rng) for name in spec.order}
    numeric = [f for f in ("quantity", "unit_price", "amount") if f in cells]
    kind = rng.choice(["no_sku", "bad_date", "bad_numbers"])
    if kind == "no_sku":
        cells["sku"] = ""
    elif kind == "bad_date":
        cells["order_date"] = rng.choice(["TBC", "see note", "??"])
    else:
        for f in numeric:
            cells[f] = rng.choice(["n/a", "-", "pending"])
    return [cells[name] for name in spec.order]


def write_messy(truth: pd.DataFrame, spec: MessSpec, path: Path, seed: int = 0, title: str = "Sales export") -> dict:
    """Write one messy sales file. Returns what a perfect reader should find in it."""
    rng = random.Random(seed)
    records = truth.to_dict("records")
    header = [spec.headers.get(name, name) for name in spec.order]
    width = len(header)

    body: list[list[object]] = [[_cell(r, name, spec, rng) for name in spec.order] for r in records]
    for _ in range(spec.duplicates):
        i = rng.randrange(len(body))
        body.insert(rng.randrange(len(body)), list(body[i]))
    for n in range(spec.bad_rows):
        body.insert(rng.randrange(len(body)), _bad_row(spec, rng.choice(records), n + 1, rng))

    # Cut the body into sheets (xlsx only), then decorate each sheet separately.
    parts = max(1, spec.split_sheets)
    size = -(-len(body) // parts)
    sheets: list[tuple[str, list[list[object]], int]] = []
    for s in range(parts):
        chunk = body[s * size:(s + 1) * size]
        rows: list[list[object]] = []
        for b in range(spec.banner_rows):
            text = [f"{title}", f"Generated {date(2026, 10, 1).strftime('%d/%m/%Y')} by ERP", "Confidential",
                    "All amounts exclude tax"][b % 4]
            rows.append([text] + [None] * (width - 1))
            if b == spec.banner_rows - 1:
                rows.append([None] * width)
        header_row = len(rows)
        rows.append(list(header))
        for i, line in enumerate(chunk):
            rows.append(line)
            if spec.repeat_header and i == len(chunk) // 2:
                rows.append(list(header))  # a page break in the original print layout
        for _ in range(spec.blank_rows):
            rows.insert(rng.randrange(header_row + 2, len(rows)), [None] * width)
        if spec.totals_row:
            total = [None] * width
            total[0] = "TOTAL"
            for f in ("amount", "quantity"):
                if f in spec.order:
                    col = spec.order.index(f)
                    if col == 0:
                        continue
                    values = truth[f].sum()
                    total[col] = format_number(round(float(values), 2), spec.number_style) if f == "amount" else int(values)
            rows.append(total)
        sheets.append((f"Batch {s + 1}" if parts > 1 else "Export", rows, header_row))

    path.parent.mkdir(parents=True, exist_ok=True)
    if spec.fmt == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows, _ in sheets:
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
            if spec.banner_rows:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(width, 4))
        if spec.extra_sheet:
            ws = wb.create_sheet("Notes")
            ws.append(["Notes for head office"])
            ws.append(["Region", "Comment"])
            ws.append(["All", "Two deliveries slipped to next month"])
        wb.save(path)
    else:
        with open(path, "w", newline="", encoding=spec.encoding, errors="replace") as fh:
            writer = csv.writer(fh, delimiter=spec.delimiter)
            for row in sheets[0][1]:
                writer.writerow(["" if c is None else c for c in row])

    mapping = {spec.headers.get(name, name): (name if name in spec.headers else None) for name in spec.order}
    return {"header_rows": {name: hr for name, _, hr in sheets}, "mapping": mapping,
            "expected_rows": len(truth), "bad_rows": spec.bad_rows, "duplicates": spec.duplicates}


# ----------------------------------------------------------------------------- the demo set

def _region_spec(region: str) -> MessSpec:
    """Three fixed layouts for the demo, one per branch, each messy in a different way."""
    if region == "North":
        return MessSpec(fmt="xlsx", banner_rows=3, date_style="dmy_slash", mixed_native_dates=True,
                        number_style="us_currency", totals_row=True, blank_rows=4, bad_rows=3, sku_mess=0.05,
                        headers={"order_id": "Order No.", "order_date": "Order Date", "customer": "Customer",
                                 "sku": "Product Code", "quantity": "Qty", "unit_price": "Unit Price",
                                 "amount": "Line Total"},
                        order=["order_id", "order_date", "customer", "sku", "quantity", "unit_price", "amount",
                               "Sales rep"])
    if region == "South":
        return MessSpec(fmt="csv", delimiter=";", encoding="cp1252", date_style="dmy_dot", number_style="eu",
                        qty_suffix=True, duplicates=6, bad_rows=2,
                        headers={"order_id": "Invoice #", "order_date": "Date", "customer": "Client",
                                 "sku": "Article", "quantity": "Quantity", "unit_price": "Price",
                                 "amount": "Amount", "region": "Branch"},
                        order=["order_date", "order_id", "region", "customer", "sku", "quantity", "unit_price",
                               "amount", "Notes"])
    return MessSpec(fmt="xlsx", banner_rows=1, date_style="mdy_slash", number_style="plain", split_sheets=4,
                    extra_sheet=True, repeat_header=True, sku_mess=0.25, totals_row=True, bad_rows=2,
                    headers={"order_id": "Order Ref", "order_date": "Booked", "customer": "Store",
                             "sku": "Item No.", "quantity": "Pieces", "unit_price": "Each",
                             "amount": "Ext. Price"},
                    order=["order_id", "order_date", "customer", "sku", "quantity", "unit_price", "amount"])


DEMO_PERIODS = ["2026-08", "2026-09"]
DEMO_VOLUMES = {"2026-08": {"North": 140, "South": 125, "West": 120},
                "2026-09": {"North": 155, "South": 105, "West": 138}}


def write_reference(folder: Path, seed: int = 7) -> None:
    """The product list (an Excel file with a banner) and the targets (a CSV, as Google Sheets exports it)."""
    folder.mkdir(parents=True, exist_ok=True)
    products = make_products(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(["Cedar & Stone Garden Supplies - product list"])
    ws.append([None])
    ws.append(["Item No.", "Description", "Family", "Cost", "List price"])
    for r in products.itertuples():
        ws.append([r.sku, r.product_name, r.category, r.unit_cost, r.list_price])
    wb.save(folder / "products.xlsx")
    targets = make_targets(DEMO_PERIODS, seed)
    with open(folder / "targets.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Region", "Month", "Target"])
        for r in targets.itertuples():
            w.writerow([r.region, r.month, r.target_revenue])


def write_demo(root: Path, seed: int = 7) -> dict[str, pd.DataFrame]:
    """Write the demo inbox for two months plus the reference files. Returns the truth per month."""
    products = make_products(seed)
    write_reference(root / "reference", seed)
    names = {"North": "sales_north_{p}.xlsx", "South": "south_export_{p}.csv", "West": "West weekly sales {p}.xlsx"}
    truths: dict[str, pd.DataFrame] = {}
    for period in DEMO_PERIODS:
        truth = make_sales(period, products, seed, DEMO_VOLUMES[period])
        # Two things every real month has: sales of a product missing from the product list
        # (they are real revenue, so they are part of the truth), and late lines from last month.
        day = truth["order_date"].iloc[0].replace(day=14)
        extra = pd.DataFrame([{"order_id": f"N{period[2:4]}{period[5:]}-9001", "order_date": day, "region": "North",
                               "sku": "HT-0199", "quantity": q, "unit_price": 14.5, "amount": round(q * 14.5, 2),
                               "customer": "Oakfield Nursery"} for q in (6, 10)])
        truth = pd.concat([truth, extra], ignore_index=True)
        truths[period] = truth
        late = truth[truth["region"] == "South"].head(3).copy()
        late["order_date"] = (day.replace(day=1) - pd.Timedelta(days=2))
        late["order_id"] = [f"S-LATE-{n}" for n in range(1, 4)]
        for i, region in enumerate(REGIONS):
            part = truth[truth["region"] == region]
            if region == "South":
                part = pd.concat([part, late], ignore_index=True)
            write_messy(part, _region_spec(region), root / "inbox" / period / names[region].format(p=period),
                        seed=seed * 100 + i, title=f"Cedar & Stone - {region} branch - sales export {period}")
        (root / "truth").mkdir(parents=True, exist_ok=True)
        truth.to_csv(root / "truth" / f"truth_{period}.csv", index=False)
    return truths
